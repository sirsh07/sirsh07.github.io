from pathlib import Path
import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt
from matplotlib.colors import ListedColormap
from matplotlib.patches import Patch
import numpy as np
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
ASSETS = ROOT / "liver_outcome_assets"
ASSETS.mkdir(exist_ok=True)
OUT = ROOT / "PFAS_liver_outcome_plots.docx"
TEAL, DARK, PURPLE, GRAY = "#087f8c", "#173b45", "#7a5195", "#607484"

def style(ax):
    ax.set_facecolor("#f9fbfb")
    ax.grid(axis="x", color="#dce7e8", linewidth=.8)
    ax.set_axisbelow(True)
    ax.spines[["top", "right"]].set_visible(False)

def save(name):
    path = ASSETS / name
    plt.savefig(path, dpi=260, bbox_inches="tight", facecolor="white")
    plt.close()
    return path

# Figure 1
names = ["AST", "GGT", "ALP"]
est, lo, hi = np.array([7.5, 9.7, 2.8]), np.array([4.0, 1.7, .5]), np.array([10.4, 17.0, 5.4])
y = np.arange(3)
fig, ax = plt.subplots(figsize=(7.4, 4.25)); style(ax)
fig.subplots_adjust(top=.80)
ax.errorbar(est, y, xerr=[est-lo, hi-est], fmt="o", markersize=8, color=TEAL, ecolor=TEAL, elinewidth=2.2, capsize=5)
ax.axvline(0, color="#768692", linewidth=1)
ax.set(yticks=y, yticklabels=names, xlabel="Percent change in liver enzyme (95% CI)", xlim=(-1,20))
ax.invert_yaxis()
fig.suptitle("PFAS mixture increase and liver-enzyme levels",x=.125,y=.97,ha="left",fontsize=14,weight="bold",color=DARK)
fig.text(.125,.875,"One-quartile simultaneous mixture increase · CHMS · n = 2,512",ha="left",color=GRAY,fontsize=9)
for yy, ee, ll, hh in zip(y, est, lo, hi):
    ax.text(hh+.35, yy, f"{ee:.1f}% ({ll:.1f}–{hh:.1f})", va="center", fontsize=9)
fig1 = save("01_liver_enzyme_change.png")

# Figure 2
names = ["PFOS → NASH", "PFHxS → NASH", "PFHxS → fibrosis", "PFHxS → lobular inflammation"]
est, lo, hi = np.array([3.32,4.18,4.44,2.87]), np.array([1.40,1.64,1.34,1.12]), np.array([7.87,10.7,14.8,7.31])
y = np.arange(4)
fig, ax = plt.subplots(figsize=(7.4, 4.6)); style(ax)
fig.subplots_adjust(top=.80)
ax.errorbar(est, y, xerr=[est-lo, hi-est], fmt="o", markersize=8, color=PURPLE, ecolor=PURPLE, elinewidth=2.2, capsize=5)
ax.axvline(1, color="#6c7a82", linestyle="--")
ax.set_xscale("log"); ax.set_xlim(.75,20); ax.set_xticks([1,2,4,8,16],["1","2","4","8","16"])
ax.set_yticks(y,names); ax.invert_yaxis(); ax.set_xlabel("Odds ratio per IQR increase in PFAS (log scale)")
fig.suptitle("PFAS exposure and clinical liver outcomes",x=.125,y=.97,ha="left",fontsize=14,weight="bold",color=DARK)
fig.text(.125,.875,"Biopsy-characterized pediatric NAFLD cohort · n = 74",ha="left",color=GRAY,fontsize=9)
for yy, ee, ll, hh in zip(y, est, lo, hi):
    ax.text(15.7,yy,f"{ee:.2f} ({ll:.2f}–{hh:.2f})",ha="right",va="center",fontsize=8.3)
fig2 = save("02_clinical_outcome_odds.png")

# Figure 3
names, vals = ["ALT","AST","GGT","Total bilirubin"], np.array([.757,.657,.461,.648])
y = np.arange(4)
fig, ax = plt.subplots(figsize=(7.4,4.25)); style(ax)
fig.subplots_adjust(top=.80)
bars = ax.barh(y,vals,color=["#1c9099","#41ae76","#78c679","#2b8cbe"],height=.62)
ax.set_yticks(y,names); ax.invert_yaxis(); ax.set_xlim(0,1)
ax.set_xticks(np.linspace(0,1,6),[f"{v:.0%}" for v in np.linspace(0,1,6)])
ax.set_xlabel("PFNA share of positive-direction WQS mixture weight")
fig.suptitle("PFNA contribution within fitted PFAS mixture models",x=.125,y=.97,ha="left",fontsize=14,weight="bold",color=DARK)
fig.text(.125,.875,"NHANES 2009–2018 · n = 7,484 · weights are not effect sizes",ha="left",color=GRAY,fontsize=9)
for bar,val in zip(bars,vals):
    ax.text(val-.025,bar.get_y()+bar.get_height()/2,f"{val:.1%}",ha="right",va="center",color="white",weight="bold")
fig3 = save("03_pfna_mixture_weights.png")

# Figure 4
studies=["Lin 2010 · NHANES","Jain 2019 · NHANES","Borghese 2022 · CHMS","Yan 2024 · NHANES","Cheng 2023 · NHANES","Jin 2020 · pediatric NAFLD","Costello 2022 · meta-analysis"]
outcomes=["ALT","AST","GGT","ALP","Bilirubin","FIB-4 / fibrosis","NAFLD / NASH"]
matrix=np.array([[1,0,1,0,0,0,0],[2,0,2,0,0,0,0],[3,1,1,1,3,0,0],[4,1,4,0,1,3,0],[0,0,0,0,0,1,3],[0,0,0,0,0,1,1],[1,2,2,0,0,0,2]])
notes=[
["PFOA ↑","","PFOA ↑","","","",""],["obese subgroup ↑","","obese subgroup ↑","","","",""],
["not consistent","mixture ↑","mixture ↑","mixture ↑","not consistent","",""],
["positive / nonlinear","mixture ↑","positive / nonlinear","","mixture ↑","p = 0.146",""],
["","","","","","PFOS / mixture ↑","individual PFAS null"],["","","","","","PFHxS ↑","PFOS/PFHxS ↑"],
["PFOA/PFOS/PFNA ↑","overall evidence","overall evidence","","","","pathway support"]]
cmap=ListedColormap(["#f4f6f7","#bee6d6","#ffe5a8","#e3e9ec","#d9c9eb"])
fig,ax=plt.subplots(figsize=(12.4,5.4)); ax.imshow(matrix,cmap=cmap,vmin=0,vmax=4,aspect="auto")
fig.subplots_adjust(top=.79,bottom=.18)
ax.set_xticks(np.arange(7),outcomes,fontsize=9); ax.set_yticks(np.arange(7),studies,fontsize=9)
ax.tick_params(top=True,bottom=False,labeltop=True,labelbottom=False,length=0)
for i in range(7):
    for j in range(7):
        ax.text(j,i,notes[i][j] or "—",ha="center",va="center",fontsize=7.1,color=DARK if matrix[i,j] else "#9aa6ab")
ax.set_xticks(np.arange(-.5,7,1),minor=True); ax.set_yticks(np.arange(-.5,7,1),minor=True)
ax.grid(which="minor",color="white",linewidth=3); ax.tick_params(which="minor",bottom=False,left=False)
for s in ax.spines.values(): s.set_visible(False)
fig.suptitle("Cross-study evidence map for PFAS and liver outcomes",x=.125,y=.97,ha="left",fontsize=15,weight="bold",color=DARK)
legend=[Patch(facecolor=c,label=l) for c,l in [("#bee6d6","Positive"),("#ffe5a8","Mixed / subgroup"),("#e3e9ec","Null / inconsistent"),("#d9c9eb","Nonlinear"),("#f4f6f7","Not reported")]]
ax.legend(handles=legend,loc="lower center",bbox_to_anchor=(.5,-.20),ncol=5,frameon=False,fontsize=8)
fig4=save("04_evidence_matrix.png")

def shade(cell, fill):
    props=cell._tc.get_or_add_tcPr(); node=props.find(qn("w:shd"))
    if node is None: node=OxmlElement("w:shd"); props.append(node)
    node.set(qn("w:fill"),fill)

def link(paragraph,text,url):
    rid=paragraph.part.relate_to(url,"http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",is_external=True)
    h=OxmlElement("w:hyperlink"); h.set(qn("r:id"),rid); r=OxmlElement("w:r"); p=OxmlElement("w:rPr")
    c=OxmlElement("w:color"); c.set(qn("w:val"),"056E83"); u=OxmlElement("w:u"); u.set(qn("w:val"),"single")
    p.extend([c,u]); r.append(p); t=OxmlElement("w:t"); t.text=text; r.append(t); h.append(r); paragraph._p.append(h)

def caption(doc,text,url):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=p.add_run(text+" "); r.italic=True; r.font.size=Pt(8.5); r.font.color.rgb=RGBColor.from_string("607484")
    link(p,"Source",url)

def callout(doc,title,text,fill="E8F5F2"):
    cell=doc.add_table(rows=1,cols=1).cell(0,0); shade(cell,fill)
    p=cell.paragraphs[0]; r=p.add_run(title+": "); r.bold=True; p.add_run(text)

doc=Document(); sec=doc.sections[0]
sec.top_margin=Inches(.7); sec.bottom_margin=Inches(.65); sec.left_margin=Inches(.72); sec.right_margin=Inches(.72)
doc.styles["Normal"].font.name="Aptos"; doc.styles["Normal"].font.size=Pt(10)
for n,s,c in [("Title",28,"173B45"),("Heading 1",18,"173B45"),("Heading 2",13,"087F8C")]:
    st=doc.styles[n]; st.font.name="Aptos Display"; st.font.size=Pt(s); st.font.bold=True; st.font.color.rgb=RGBColor.from_string(c)
header=sec.header.paragraphs[0]; header.text="PFAS EXPOSURE AND LIVER OUTCOMES"; header.alignment=WD_ALIGN_PARAGRAPH.RIGHT
footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=footer.add_run()
b=OxmlElement("w:fldChar"); b.set(qn("w:fldCharType"),"begin"); i=OxmlElement("w:instrText"); i.text=" PAGE "; e=OxmlElement("w:fldChar"); e.set(qn("w:fldCharType"),"end"); r._r.extend([b,i,e])

p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.space_before=Pt(120)
r=p.add_run("PFAS Exposure and Liver Outcomes"); r.bold=True; r.font.size=Pt(29); r.font.color.rgb=RGBColor.from_string("173B45")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run("Plots, interpretation, evidence summary, and research implications"); r.font.size=Pt(15); r.font.color.rgb=RGBColor.from_string("087F8C")
p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; p.add_run("Prepared 23 September 2026").bold=True
doc.add_page_break()

doc.add_heading("How to read these plots",1)
doc.add_paragraph("The figures summarize adjusted estimates from published human studies. Percent changes, odds ratios, regression coefficients, meta-analytic z-scores, and mixture weights answer different questions and are therefore kept separate.")
callout(doc,"Interpretation guardrail","These observational associations do not alone prove that PFAS caused liver injury. Population, exposure range, year, assay, outcome definition, and statistical adjustment differ across studies.","FFF2D9")
doc.add_heading("Main findings",2)
for x in ["ALT, AST, and GGT show the most recurring positive signals.","Mixture analyses find associations even when individual-compound results vary.","Pediatric PFOS/PFHxS estimates point toward NASH and fibrosis, but come from a small cohort.","NAFLD and FIB-4 findings are less uniform and can be null, subgroup-specific, or nonlinear."]:
    doc.add_paragraph(x,style="List Bullet")

sections=[
("Liver-enzyme change",fig1,"Figure 1. Adjusted percent change per one-quartile simultaneous increase in the PFAS mixture; bars are 95% CIs.","https://pmc.ncbi.nlm.nih.gov/articles/PMC9472375/","AST, GGT, and ALP confidence intervals are above zero. GGT has the largest point estimate but the widest interval. ALT and bilirubin were not consistently associated in this analysis."),
("Clinical liver outcomes",fig2,"Figure 2. Adjusted odds ratios per IQR increase in PFOS or PFHxS; bars are 95% CIs.","https://pmc.ncbi.nlm.nih.gov/articles/PMC6944061/","The estimates point toward higher odds of NASH, fibrosis, and lobular inflammation. Wide intervals reflect the cohort size of 74, so these are suggestive rather than precise population-wide risk estimates."),
("Mixture contribution",fig3,"Figure 3. PFNA weight in positive-direction WQS models for four liver biomarkers.","https://pmc.ncbi.nlm.nih.gov/articles/PMC11629162/","PFNA received the largest reported contribution for these outcomes. WQS weights describe mixture composition; they are not percentages of disease caused or effect sizes."),
("Consistency across studies",fig4,"Figure 4. Qualitative map of adjusted findings across selected studies.","https://pmc.ncbi.nlm.nih.gov/articles/PMC9044977/","The map shows recurring enzyme signals and substantial heterogeneity. It also shows why a single correlation between regional PFAS summaries and regional liver outcomes would be misleading.")]
for heading,image,cap,url,body in sections:
    doc.add_heading(heading,1); doc.add_picture(str(image),width=Inches(6.7)); caption(doc,cap,url); doc.add_paragraph(body)

doc.add_heading("Selected estimates using other scales",1)
table=doc.add_table(rows=1,cols=5); table.style="Table Grid"; table.alignment=WD_TABLE_ALIGNMENT.CENTER
for c,v in zip(table.rows[0].cells,["Study","Exposure","Outcome","Adjusted estimate","Population"]):
    shade(c,"173B45"); r=c.paragraphs[0].add_run(v); r.bold=True; r.font.color.rgb=RGBColor(255,255,255); r.font.size=Pt(8)
rows=[
["Lin 2010","PFOA","ALT","+1.86 U/L (1.24–2.48)","NHANES; n=2,216"],["Lin 2010","PFOA","log GGT","β=.08 (.05–.11)","NHANES; n=2,216"],
["Cheng 2023","PFOS","FIB-4","β=.07 (.01–.13)","NHANES; n=1,150"],["Jin 2020","PFHxS","NAFLD activity","β=.46 (.03–.89)","Pediatric; n=74"],
["Costello 2022","PFOA","ALT","weighted z=6.20","24 studies"],["Costello 2022","PFOS","ALT","weighted z=3.55","24 studies"],["Costello 2022","PFNA","ALT","weighted z=2.27","24 studies"]]
for n,row in enumerate(rows):
    cells=table.add_row().cells
    if n%2:
        for c in cells: shade(c,"F1F6F7")
    for c,v in zip(cells,row): c.text=v; c.paragraphs[0].runs[0].font.size=Pt(7.8)
doc.add_paragraph("These estimates are tabulated rather than placed on one axis because their units and statistical meanings are not commensurate.")

doc.add_heading("Recommended participant-level analysis",1)
doc.add_paragraph("Plot log₂ serum PFAS against log-transformed ALT, AST, and GGT for the same participants, using a flexible smoother with a 95% confidence band. Then report adjusted estimates per exposure doubling or per IQR.")
for title,body in [
("Describe","Examine distributions, detection frequencies, missingness, and PFAS correlations."),
("Visualize","Use individual-level scatter or hexbin plots with nonlinear smoothers."),
("Adjust","Include age, sex, race/ethnicity, BMI or waist, alcohol, smoking, diabetes, lipids, medications, kidney function, fasting, hepatitis status, and survey design where applicable."),
("Model mixtures","Compare individual-PFAS models with a prespecified mixture method."),
("Validate","Report sensitivity analyses, multiplicity control, calibration, and external validation where possible.")]:
    p=doc.add_paragraph(); p.add_run(title+" — ").bold=True; p.add_run(body)
callout(doc,"Primary endpoints","Use ALT, AST, and GGT as primary continuous outcomes. Treat ALP, bilirubin, FIB-4, and diagnosed NAFLD/NASH as secondary unless sample size and clinical ascertainment are strong.")

doc.add_heading("References",1)
refs=[
("Borghese et al. 2022 — CHMS mixture analysis","https://pmc.ncbi.nlm.nih.gov/articles/PMC9472375/"),
("Jin et al. 2020 — pediatric NAFLD","https://pmc.ncbi.nlm.nih.gov/articles/PMC6944061/"),
("Lin et al. 2010 — NHANES liver biomarkers","https://pubmed.ncbi.nlm.nih.gov/20010922/"),
("Jain and Ducatman 2019 — obesity-stratified NHANES","https://pubmed.ncbi.nlm.nih.gov/30589657/"),
("Cheng et al. 2023 — NHANES NAFLD and fibrosis","https://pmc.ncbi.nlm.nih.gov/articles/PMC10281433/"),
("Yan et al. 2024 — NHANES mixture analysis","https://pmc.ncbi.nlm.nih.gov/articles/PMC11629162/"),
("Costello et al. 2022 — systematic review and meta-analysis","https://pmc.ncbi.nlm.nih.gov/articles/PMC9044977/")]
for title,url in refs:
    p=doc.add_paragraph(style="List Number"); link(p,title,url)
doc.add_paragraph("The companion PFAS_liver_outcome_effects_data.csv contains the numeric values, study details, confidence intervals, interpretations, and source URLs.")
doc.save(OUT)
print("Created",OUT)
for x in [fig1,fig2,fig3,fig4]: print("Created",x)
