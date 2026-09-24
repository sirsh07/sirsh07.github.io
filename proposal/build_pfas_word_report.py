from pathlib import Path
from datetime import date

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION, WD_ORIENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parent
OUT = ROOT / "PFAS_datasets_evidence_and_proposed_study.docx"
ASSETS = ROOT / "report_assets"

TEAL = "087F75"
DARK = "17343D"
PALE_TEAL = "DFF3EF"
PALE_BLUE = "E7F0F7"
PALE_AMBER = "FFF0D9"
GRAY = "5F7078"
WHITE = "FFFFFF"


def shade(cell, color):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), color)


def margins(section, top=.65, bottom=.65, left=.68, right=.68):
    section.top_margin = Inches(top)
    section.bottom_margin = Inches(bottom)
    section.left_margin = Inches(left)
    section.right_margin = Inches(right)


def set_cell_text(cell, text, bold=False, color=None, size=8.2):
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(str(text))
    r.bold = bold
    r.font.size = Pt(size)
    if color:
        r.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return p


def add_hyperlink(paragraph, text, url, color="24577D"):
    part = paragraph.part
    rid = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color")
    c.set(qn("w:val"), color)
    rpr.append(c)
    u = OxmlElement("w:u")
    u.set(qn("w:val"), "single")
    rpr.append(u)
    run.append(rpr)
    t = OxmlElement("w:t")
    t.text = text
    run.append(t)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)
    return hyperlink


def add_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_begin = OxmlElement("w:fldChar")
    fld_begin.set(qn("w:fldCharType"), "begin")
    instr = OxmlElement("w:instrText")
    instr.set(qn("xml:space"), "preserve")
    instr.text = " PAGE "
    fld_end = OxmlElement("w:fldChar")
    fld_end.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_begin, instr, fld_end])


def add_caption(doc, text, source_url=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.italic = True
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(GRAY)
    if source_url:
        p.add_run(" ")
        add_hyperlink(p, "Source", source_url)
    return p


def add_callout(doc, title, text, color=PALE_TEAL):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade(cell, color)
    p = cell.paragraphs[0]
    r = p.add_run(title + ": ")
    r.bold = True
    r.font.color.rgb = RGBColor.from_string(DARK)
    p.add_run(text)
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def add_flow(doc, nodes, caption):
    if len(nodes) >= 5:
        table = doc.add_table(rows=len(nodes) * 2 - 1, cols=1)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        for idx, node in enumerate(nodes):
            row = idx * 2
            cell = table.cell(row, 0)
            shade(cell, PALE_TEAL if idx % 2 == 0 else PALE_BLUE)
            p = set_cell_text(cell, "", size=9)
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            title, body, citation = node
            r = p.add_run(title + "\n")
            r.bold = True
            r.font.color.rgb = RGBColor.from_string(TEAL)
            r.font.size = Pt(10)
            rr = p.add_run(body + "\n")
            rr.font.size = Pt(8.5)
            rc = p.add_run(citation)
            rc.italic = True
            rc.font.size = Pt(7.5)
            if idx < len(nodes) - 1:
                arrow = table.cell(row + 1, 0)
                p2 = set_cell_text(arrow, "↓", bold=True, color=TEAL, size=16)
                p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_caption(doc, caption)
        return
    table = doc.add_table(rows=1, cols=len(nodes) * 2 - 1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for idx, node in enumerate(nodes):
        col = idx * 2
        cell = table.cell(0, col)
        shade(cell, PALE_TEAL if idx % 2 == 0 else PALE_BLUE)
        p = set_cell_text(cell, "", size=9)
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title, body, citation = node
        r = p.add_run(title + "\n")
        r.bold = True
        r.font.color.rgb = RGBColor.from_string(TEAL)
        r.font.size = Pt(10)
        rr = p.add_run(body + "\n")
        rr.font.size = Pt(8.3)
        rc = p.add_run(citation)
        rc.italic = True
        rc.font.size = Pt(7.5)
        if idx < len(nodes) - 1:
            arrow = table.cell(0, col + 1)
            p2 = set_cell_text(arrow, "→", bold=True, color=TEAL, size=18)
            p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_caption(doc, caption)


doc = Document()
margins(doc.sections[0])

# Global styles
styles = doc.styles
styles["Normal"].font.name = "Aptos"
styles["Normal"].font.size = Pt(10)
styles["Normal"].paragraph_format.space_after = Pt(5)
for style_name, size, color in [("Title", 28, DARK), ("Heading 1", 18, DARK), ("Heading 2", 13, TEAL), ("Heading 3", 11, DARK)]:
    st = styles[style_name]
    st.font.name = "Aptos Display"
    st.font.size = Pt(size)
    st.font.color.rgb = RGBColor.from_string(color)
    st.font.bold = True
styles["Heading 1"].paragraph_format.space_before = Pt(14)
styles["Heading 1"].paragraph_format.space_after = Pt(6)
styles["Heading 2"].paragraph_format.space_before = Pt(10)
styles["Heading 2"].paragraph_format.space_after = Pt(4)

# Header/footer
header = doc.sections[0].header.paragraphs[0]
header.text = "PFAS DATASETS AND EVIDENCE BRIEF"
header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
header.runs[0].font.size = Pt(8)
header.runs[0].font.color.rgb = RGBColor.from_string(GRAY)
add_page_number(doc.sections[0].footer.paragraphs[0])

# Title page
p = doc.add_paragraph()
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.space_before = Pt(105)
r = p.add_run("PFAS, Liver Health, and Lipid Metabolism")
r.bold = True
r.font.name = "Aptos Display"
r.font.size = Pt(28)
r.font.color.rgb = RGBColor.from_string(DARK)
p2 = doc.add_paragraph()
p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = p2.add_run("Datasets, prior evidence, figures, outcomes, and proposed national-to-local study design")
r.font.size = Pt(15)
r.font.color.rgb = RGBColor.from_string(TEAL)
doc.add_paragraph()
p3 = doc.add_paragraph()
p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
p3.add_run("Prepared 23 September 2026\n").bold = True
p3.add_run("Companion document to the interactive PFAS Dataset Explorer")
p4 = doc.add_paragraph()
p4.alignment = WD_ALIGN_PARAGRAPH.CENTER
add_hyperlink(p4, "Open the local interactive explorer", "pfas_dataset_explorer/index.html")
doc.add_page_break()

# Executive summary
doc.add_heading("Executive summary", level=1)
doc.add_paragraph(
    "Per- and polyfluoroalkyl substances (PFAS) are persistent environmental contaminants with recurring epidemiologic associations with liver enzymes and circulating lipids. The most reproducible human findings involve alanine aminotransferase (ALT), aspartate aminotransferase (AST), gamma-glutamyl transferase (GGT), total cholesterol, and LDL/non-HDL cholesterol. Findings for alkaline phosphatase (ALP), bilirubin, HDL, triglycerides, steatosis, fibrosis, and liver cancer are less consistent or depend strongly on compound, population, endpoint definition, and model."
)
doc.add_paragraph(
    "The central proposal opportunity is therefore not to demonstrate for the first time that PFAS can be associated with liver or lipid biomarkers. That has already been studied using NHANES, the Canadian Health Measures Survey, exposed-community cohorts, mixture models, splines, and meta-analysis. The stronger contribution is to determine whether relationships learned in nationally representative NHANES data transport to an independently characterized PFAS-exposed Florida population, to quantify calibration failure, and to develop a prespecified recalibration or adaptation strategy."
)
add_callout(doc, "Recommended design", "Develop and internally validate models in harmonized NHANES 2003–2018; freeze all model choices; test temporal performance in NHANES 2021–2023; then conduct true external validation and recalibration in the Florida cohort. Use EPA and Florida drinking-water data as environmental context rather than substitutes for serum PFAS.")

doc.add_heading("Primary objectives", level=2)
for text in [
    "Estimate individual and mixture associations of serum PFAS with prespecified liver and lipid endpoints in a harmonized national sample.",
    "Evaluate nonlinear exposure–response relationships and effect modification without presenting exploratory subgroup findings as confirmatory.",
    "Test discrimination, prediction error, and—most importantly—calibration in an independent Florida exposed-community cohort.",
    "Identify why transport fails, including differences in PFAS mixtures, exposure ranges, outcome measurement, demographics, recruitment, and local water sources.",
    "Evaluate transparent recalibration or limited-data adaptation methods while preserving an untouched external test set."
]:
    doc.add_paragraph(text, style="List Bullet")

doc.add_heading("Conceptual framework", level=2)
add_flow(doc, [
    ("Environmental sources", "Drinking water, food, dust, consumer products, and occupational sources", "EPA UCMR 5; USGS; Florida DEP"),
    ("Internal exposure", "Serum concentrations of legacy and emerging PFAS, individually and as mixtures", "NHANES; CHMS; GenX"),
    ("Hepatic response", "Lipid, bile-acid, amino-acid, glucose, inflammatory, and oxidative pathways", "Costello 2022; Goodrich 2022"),
    ("Observable outcomes", "Liver enzymes, bilirubin, fibrosis indices, cholesterol fractions, and triglycerides", "Borghese 2022; Liu 2023"),
], "Figure 1. Conceptual pathway from PFAS sources to measurable liver and lipid outcomes. Arrows represent hypothesized pathways, not proof of causality.")

# Evidence narrative
doc.add_heading("What prior research establishes", level=1)
doc.add_heading("Liver enzymes and liver injury", level=2)
doc.add_paragraph(
    "Early NHANES analyses reported positive associations of PFOA with ALT and GGT, with some evidence that obesity, insulin resistance, or metabolic syndrome strengthened associations (Lin et al., 2010). NHANES 2007–2010 analyses subsequently reported compound-specific relationships across ALT, AST, GGT, ALP, bilirubin, and uric acid (Gleason et al., 2015). In NHANES 2011–2014, associations of PFOA, PFHxS, and PFNA with ALT and of PFOA/PFNA with GGT were concentrated among participants with obesity (Jain and Ducatman, 2019)."
)
doc.add_paragraph(
    "Mixture analyses reinforce the signal while showing endpoint specificity. In the Canadian Health Measures Survey, a simultaneous one-quartile mixture increase was associated with higher AST, GGT, and ALP, but not consistently with ALT or bilirubin (Borghese et al., 2022). In NHANES 2009–2018, weighted quantile sum and spline analyses found positive mixture associations with several liver biomarkers and nonlinear relationships for selected compounds; the mixture estimate for FIB-4 was not statistically significant (Yan et al., 2024). A systematic review of 24 epidemiologic and 85 rodent studies concluded that ALT was positively associated with PFOA, PFOS, and PFNA in human evidence, while experimental evidence consistently supported hepatic steatosis and injury (Costello et al., 2022)."
)

doc.add_heading("Lipids and lipoproteins", level=2)
doc.add_paragraph(
    "The lipid evidence is strongest for total cholesterol and LDL/non-HDL cholesterol. In 46,294 adults from the highly exposed C8 Health Project, total cholesterol increased approximately 11–12 mg/dL from the lowest to highest PFOA or PFOS decile, while HDL showed no association (Steenland et al., 2009). Similar positive, nonlinear associations with total and LDL cholesterol were reported in 12,476 children and adolescents from the same project (Frisbee et al., 2010). In NHANES 2003–2004, the highest versus lowest exposure quartiles corresponded to 13.4 mg/dL higher total cholesterol for PFOS, 9.8 mg/dL for PFOA, and 13.9 mg/dL for PFNA, although PFHxS was inversely associated (Nelson et al., 2010)."
)
doc.add_paragraph(
    "Longitudinal evidence strengthens—but does not settle—causal interpretation. Declines in PFOA and PFOS in a C8 follow-up sample were accompanied by declines in total and LDL cholesterol (Fitz-Simon et al., 2013). In the Swedish PIVUS cohort, changes in six of eight PFAS were positively associated with changes in plasma lipids across three measurements over ten years. Detailed NMR profiling in the POEM cohort found that PFOS, PFOA, and PFDA were positively associated with cholesterol across multiple IDL and LDL fractions, whereas triglyceride findings were mostly null (Haug et al., 2023). A 29-study meta-analysis concluded that PFOA and PFOS were positively associated with selected lipid measures, particularly total cholesterol and LDL, while emphasizing substantial heterogeneity (Liu et al., 2023)."
)

doc.add_heading("Steatosis, fibrosis, and clinical disease", level=2)
doc.add_paragraph(
    "Evidence for clinically defined steatotic liver disease is less stable than evidence for enzymes. NHANES analyses have produced conclusions that vary with the selected fatty-liver index, survey period, and use of transient elastography (Cheng et al., 2023; Momo et al., 2024). Because common indices include BMI, triglycerides, or GGT, their association with PFAS can be difficult to separate from the components used to construct the outcome."
)
doc.add_paragraph(
    "Small studies with stronger clinical phenotyping provide suggestive signals. Among 74 children with biopsy-confirmed NAFLD, per-IQR increases in PFOS and PFHxS were associated with NASH, and PFHxS was associated with fibrosis, lobular inflammation, and higher disease activity (Jin et al., 2020). A nested analysis of 50 incident non-viral hepatocellular carcinoma cases and 50 controls reported an association with high pre-diagnostic PFOS and implicated altered glucose, amino-acid, and bile-acid pathways (Goodrich et al., 2022). However, a newer individual-participant analysis across 12 prospective cohorts did not establish a clear overall association with liver cancer, so this remains an exploratory endpoint rather than a central proposal claim."
)

# Dataset table landscape
land = doc.add_section(WD_SECTION.NEW_PAGE)
land.orientation = WD_ORIENT.LANDSCAPE
land.page_width, land.page_height = land.page_height, land.page_width
margins(land, .55, .55, .45, .45)
add_page_number(land.footer.paragraphs[0])
doc.add_heading("Datasets that can support the research", level=1)
doc.add_paragraph("Counts below distinguish people from environmental samples or systems. Published analytic samples are usually smaller than released source-file counts because of age restrictions, missing outcomes, covariate completeness, detection rules, fasting requirements, and repeated observations.")

datasets = [
    ("NHANES", "U.S.; national", "1999–2000; 2003–2023", "18,236 PFAS-file rows, 2003–2018; 23,181 non-overlapping rows, 2003–2023, before exclusions", "Serum PFAS; liver chemistry; lipids; platelets; extensive covariates", "Open XPT files", "Primary development and temporal validation", "https://wwwn.cdc.gov/Nchs/Nhanes/Search/DataPage.aspx?Component=Laboratory"),
    ("Canadian Health Measures Survey", "Canada; national", "2007–2011; 2016–2017 in cited analysis", "n=1,957–4,657 depending PFAS/outcome", "Six PFAS; ALT, AST, GGT, ALP, bilirubin; survey design", "Controlled RDC", "Independent national replication", "https://www.statcan.gc.ca/en/microdata/data-centres/access"),
    ("POEM", "Uppsala, Sweden", "2010–2016", "n=493; all age 50", "Four PFAS; 43 NMR lipid/lipoprotein measures", "Author request", "Detailed lipid validation", "https://pmc.ncbi.nlm.nih.gov/articles/PMC10541331/"),
    ("Pittsboro paired cohort", "Pittsboro, North Carolina", "2019–2020", "49 adults; 92 paired observations; 43 at both visits", "13 PFAS in water/serum; lipid panel; metabolic panel", "Investigator/IRB request", "Closest small community analogue", "https://pmc.ncbi.nlm.nih.gov/articles/PMC10529814/"),
    ("GenX Exposure Study", "Cape Fear River Basin, NC", "2017 onward", ">1,000 overall; 1,020 blood samples in 2020–2021; wave-specific analytic n", "Legacy/emerging PFAS; water; serum; questionnaires; selected clinical outcomes", "Investigator request", "Exposed-community external comparison", "https://genxstudy.ncsu.edu/"),
    ("ATSDR 10-site assessments", "Ten U.S. exposed communities", "Approx. 2018–2021", "2,384 residents; 1,212 households", "Seven common serum PFAS; water/exposure history; demographics", "Reports open; microdata not direct", "Exposure-distribution benchmark", "https://www.atsdr.cdc.gov/pfas/final-report/index.html"),
    ("MIREC", "Ten Canadian cities", "2008 onward", "Approx. 2,000 original pregnancies", "PFAS and other chemicals; maternal/child metabolic outcomes", "Biobank application", "Life-course extension", "https://www.canada.ca/en/health-canada/services/environmental-workplace-health/environmental-contaminants/human-biomonitoring-environmental-chemicals/maternal-infant-research-environmental-chemicals-mirec-study.html"),
    ("EPA UCMR 5", "United States", "2023–2025", "~10,130 systems; ~25,950 sampling locations", "29 PFAS plus lithium; no individual outcomes", "Open text/Excel", "National water-occurrence context", "https://www.epa.gov/dwucmr/fifth-unregulated-contaminant-monitoring-rule-data-finder"),
    ("USGS national tap water", "U.S., including AK, HI, PR", "2021–2022", "409 one-time + 85 repeated samples", "34 PFAS; no individual outcomes", "Open download", "Independent water-occurrence mapping", "https://www.usgs.gov/data/concentrations-and-polyfluoroalkyl-substances-pfas-tapwater-collected-throughout-united-states"),
    ("EPA PFAS Analytic Tools", "United States", "Multiple years", "Dynamic records", "Integrated water, facilities, releases, waste, tissue and other media", "Open filtered exports", "Source and spatial overlays", "https://echo.epa.gov/trends/pfas-tools"),
    ("Florida DEP drinking-water data", "Florida", "Historical/current; monitoring expanding through 2027", "Dynamic system/sample records", "Public-water-system, source, service-population and chemistry information", "Open reports/data", "Florida-specific environmental context", "https://floridadep.gov/water/source-drinking-water/content/chemical-data"),
]

headers = ["Dataset", "Area", "Years", "Size", "Relevant contents", "Access", "Best use", "Source"]
table = doc.add_table(rows=1, cols=len(headers))
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.style = "Table Grid"
for i, h in enumerate(headers):
    set_cell_text(table.rows[0].cells[i], h, bold=True, color=WHITE, size=8)
    shade(table.rows[0].cells[i], TEAL)
for row in datasets:
    cells = table.add_row().cells
    for i, val in enumerate(row[:-1]):
        set_cell_text(cells[i], val, bold=(i == 0), size=7.4)
        if len(table.rows) % 2 == 0:
            shade(cells[i], "F2F7F5")
    p = set_cell_text(cells[-1], "", size=7.2)
    add_hyperlink(p, "Open", row[-1])

doc.add_paragraph()
add_callout(doc, "Dataset decision", "NHANES is the only major source here that is immediately downloadable and jointly contains person-level serum PFAS, liver outcomes, lipid outcomes, covariates, and survey weights at national scale. The Florida cohort is the required independent external test. Environmental datasets improve exposure context but cannot train the health-outcome model alone.")

# Back to portrait
portrait = doc.add_section(WD_SECTION.NEW_PAGE)
portrait.orientation = WD_ORIENT.PORTRAIT
portrait.page_width, portrait.page_height = portrait.page_height, portrait.page_width
margins(portrait)
add_page_number(portrait.footer.paragraphs[0])

doc.add_heading("Recommended dataset architecture", level=1)
add_flow(doc, [
    ("Development", "Harmonized NHANES 2003–2018; cycle-aware resampling; prespecified features and outcomes", "18,236 PFAS-file rows before exclusions"),
    ("Temporal validation", "Lock the model, then test in NHANES 2021–2023 without further tuning", "3,618 PFAS-file rows before exclusions"),
    ("External validation", "Apply unchanged model to the Florida serum cohort; report calibration and error", "Independent exposed population"),
    ("Adaptation", "If transport fails, recalibrate intercept/slope or use limited local updating with separate testing", "Preserve an untouched evaluation set"),
], "Figure 2. Recommended national-to-local model-development and validation sequence.")

doc.add_heading("Variables needed from the Florida cohort", level=2)
for text in [
    "Participant count, recruitment years, counties and water systems, eligibility criteria, repeat visits, and specimen timing.",
    "PFAS analyte list, matrix, units, assay method, LOD/LLOQ, detection frequencies, isomer reporting, batch identifiers, and storage conditions.",
    "ALT, AST, GGT, ALP, bilirubin, platelet count, total cholesterol, HDL, LDL, triglycerides, fasting status, medication use, and laboratory reference ranges.",
    "Age, sex, race/ethnicity, BMI, smoking, alcohol, diabetes, kidney function, socioeconomic variables, pregnancy/menopause where relevant, and recruitment structure.",
    "A blinded feasibility count for complete overlap of exposure, outcome, and covariates. This—not total enrollment—is the effective external-validation sample."
]:
    doc.add_paragraph(text, style="List Bullet")

# Prior research table
doc.add_heading("Selected prior studies and conclusions", level=1)
studies = [
    ("Lin et al., 2010", "NHANES 1999–2000/2003–2004; n=2,216", "PFOA positively associated with ALT and GGT; modification by metabolic status suggested."),
    ("Gleason et al., 2015", "NHANES 2007–2010; n=4,333", "Compound-specific associations across ALT, AST, GGT, ALP and bilirubin."),
    ("Jain & Ducatman, 2019", "NHANES 2011–2014; n=2,883 across BMI strata", "ALT/GGT associations were concentrated among participants with obesity."),
    ("Borghese et al., 2022", "CHMS; n=1,957–4,657", "PFAS mixture associated with AST, GGT and ALP; ALT/bilirubin not consistent."),
    ("Yan et al., 2024", "NHANES 2009–2018; n=7,484", "Positive mixture signals for several biomarkers; nonlinear relationships; FIB-4 mixture result not significant."),
    ("Costello et al., 2022", "24 epidemiologic + 85 rodent studies", "Human ALT signal for PFOA, PFOS and PFNA; strong experimental liver-injury/steatosis evidence."),
    ("Jin et al., 2020", "74 children with biopsy-confirmed NAFLD", "PFOS/PFHxS associated with NASH; PFHxS associated with fibrosis and inflammation."),
    ("Cheng et al., 2023", "NHANES 2017–2018; n=1,150", "Stronger evidence for fibrosis indices than steatosis; no significant adjusted individual-PFAS NAFLD result."),
    ("Momo et al., 2024", "NHANES 2003–2018; n=10,234", "NAFLD associations varied by HSI, FLI and elastography definitions."),
    ("Goodrich et al., 2022", "50 HCC cases + 50 controls", "High pre-diagnostic PFOS associated with HCC in a small proof-of-concept analysis."),
    ("Steenland et al., 2009", "C8 adults; n=46,294", "PFOA/PFOS associated with total and other atherogenic lipids; HDL null."),
    ("Frisbee et al., 2010", "C8 youth; n=12,476", "Positive, nonlinear PFOA/PFOS relationships with total and LDL cholesterol."),
    ("Nelson et al., 2010", "NHANES 2003–2004; n=860 for main cholesterol outcomes", "Positive TC associations for PFOS, PFOA and PFNA; inverse PFHxS finding."),
    ("Fitz-Simon et al., 2013", "C8 follow-up; n=560", "Declining PFOA/PFOS tracked with declining total and LDL cholesterol."),
    ("PIVUS, 2022", "Sweden; n=864 baseline, repeated over 10 years", "Changes in six of eight PFAS positively associated with changes in plasma lipids."),
    ("Haug et al., 2023", "POEM; n=493", "Positive cholesterol-subfraction associations for PFOS/PFOA/PFDA; triglycerides mostly null."),
    ("GenX lipid study, 2022", "Wilmington, NC; n=326", "PFOS and PFNA associated with higher total and non-HDL cholesterol."),
    ("Liu et al., 2023", "29-study adult lipid meta-analysis", "Positive PFOA/PFOS associations, especially TC and LDL; substantial heterogeneity."),
]
st = doc.add_table(rows=1, cols=3)
st.style = "Table Grid"
st.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(["Study", "Population / size", "Principal conclusion"]):
    set_cell_text(st.rows[0].cells[i], h, True, WHITE, 8.5); shade(st.rows[0].cells[i], TEAL)
for i, row in enumerate(studies):
    cells = st.add_row().cells
    for j, val in enumerate(row):
        set_cell_text(cells[j], val, bold=(j == 0), size=8)
        if i % 2: shade(cells[j], "F2F7F5")

# Figures
doc.add_page_break()
doc.add_heading("Key published figures", level=1)
doc.add_paragraph("The figures below are reproduced from openly accessible articles and are accompanied by proposal-focused interpretation. They should be read with each paper's population, design, model, and uncertainty in mind.")

figures = [
    ("borghese_2022_fig1.png", "Figure 3. Individual and mixture associations of PFAS with AST, GGT, and ALP in the Canadian Health Measures Survey. The mixture result supports multi-pollutant modeling but does not by itself establish causation.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC9472375/", 6.7),
    ("yan_2024_fig5.jpg", "Figure 4. Weighted quantile sum contributions for five PFAS and liver outcomes in NHANES 2009–2018. A larger weight indicates contribution within the constrained mixture model, not unique causal potency.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC11629162/", 6.7),
    ("yan_2024_fig4.jpg", "Figure 5. Restricted cubic spline relationships between PFAS and liver biomarkers. Several curves are nonlinear, and uncertainty increases at exposure extremes.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC11629162/", 6.7),
    ("haug_2023_fig2.png", "Figure 6. PFAS associations with cholesterol across lipoprotein subfractions. PFOS, PFOA, and PFDA showed broad positive patterns, particularly across IDL and LDL fractions.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC10541331/", 6.7),
    ("haug_2023_fig3.png", "Figure 7. PFAS associations with triglycerides across lipoprotein subfractions. Results were largely null, illustrating why triglycerides should not be assumed to behave like cholesterol.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC10541331/", 6.7),
]
for filename, caption, url, width in figures:
    path = ASSETS / filename
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(str(path), width=Inches(width))
    add_caption(doc, caption, url)
    doc.add_paragraph()

# Outcome synthesis
doc.add_heading("Outcome-level synthesis", level=1)
outcomes = [
    ("ALT", "Moderately strong", "Usually positive for PFOA/PFOS/PFNA, but not uniform; obesity and nonlinearity may modify effects.", "Primary"),
    ("AST", "Moderate", "Positive in several individual and mixture analyses; less liver-specific than ALT.", "Primary"),
    ("GGT", "Moderately strong", "One of the more reproducible enzyme outcomes; alcohol and metabolic status require careful control.", "Primary"),
    ("ALP", "Moderate/limited", "Positive mixture findings but fewer studies and limited liver specificity.", "Primary/secondary"),
    ("Bilirubin", "Limited/mixed", "Compound- and study-specific results; use as a secondary endpoint.", "Secondary"),
    ("Total cholesterol", "Strongest lipid signal", "Positive across national, highly exposed, cross-sectional, and longitudinal studies.", "Primary"),
    ("LDL / non-HDL", "Strong", "Generally positive; non-HDL is useful in non-fasting cohorts.", "Primary"),
    ("HDL", "Mixed", "Positive, inverse, or null depending compound and population; concentration may miss particle composition.", "Secondary"),
    ("Triglycerides", "Mixed", "Frequently null or heterogeneous; fasting and subfraction definitions matter.", "Secondary/exploratory"),
    ("Steatosis / MASLD", "Inconsistent", "Sensitive to case definition, analytic cycle, BMI, and index construction.", "Exploratory"),
    ("Fibrosis / NASH", "Emerging", "Suggestive pediatric biopsy and selected index signals; limited certainty.", "Secondary/exploratory"),
    ("Liver cancer", "Early/inconclusive", "Small positive PFOS study not yet confirmed by broader pooled evidence.", "Not primary"),
]
ot = doc.add_table(rows=1, cols=4)
ot.style = "Table Grid"
ot.alignment = WD_TABLE_ALIGNMENT.CENTER
for i, h in enumerate(["Outcome", "Evidence pattern", "Interpretation", "Proposal role"]):
    set_cell_text(ot.rows[0].cells[i], h, True, WHITE, 8.5); shade(ot.rows[0].cells[i], TEAL)
for i, row in enumerate(outcomes):
    cells = ot.add_row().cells
    for j, val in enumerate(row):
        set_cell_text(cells[j], val, bold=(j in (0, 1)), size=8)
        if i % 2: shade(cells[j], "F2F7F5")

doc.add_heading("Evidence-to-conclusion flowchart", level=2)
add_flow(doc, [
    ("Repeated biomonitoring signal", "Common PFAS are detectable in national and exposed-community serum samples.", "NHANES; CHMS; C8; GenX"),
    ("Biomarker associations", "Liver enzymes and atherogenic cholesterol fractions show the most reproducible human associations.", "Costello 2022; Liu 2023"),
    ("Clinical endpoints less settled", "Steatosis, fibrosis and cancer results depend on endpoint definition and smaller samples.", "Jin 2020; Momo 2024; pooled HCC study"),
    ("Research gap", "National associations may not remain calibrated in a locally exposed community with a different mixture.", "Transportability question"),
    ("Proposed contribution", "Quantify external validity, diagnose failure, and recalibrate using limited Florida data.", "National-to-local modeling"),
], "Figure 8. Synthesis of prior evidence and the resulting proposal contribution.")

doc.add_heading("Most important findings", level=1)
for text in [
    "Total cholesterol and LDL/non-HDL cholesterol are the most consistent lipid outcomes across general-population, highly exposed, and longitudinal studies.",
    "ALT, AST, and GGT are the most defensible primary liver outcomes; ALP remains useful but is less liver-specific, and bilirubin is less consistent.",
    "Mixture and nonlinear analyses are scientifically justified, but neither is novel by itself because both have already been used in NHANES and CHMS.",
    "NAFLD/MASLD findings are definition-sensitive. Imaging or biopsy endpoints are preferable to indices that reuse BMI, triglycerides, or GGT when possible.",
    "The largest unresolved methodological problem is transportability: whether national PFAS–outcome functions remain accurate in an exposed Florida population.",
    "Prediction performance must include calibration, not only R², AUC, or mean error. Poor calibration is itself an informative finding about population shift.",
    "Environmental water data can explain exposure context and domain shift but cannot substitute for linked person-level serum PFAS and clinical outcomes.",
    "All primary hypotheses, exposure definitions, endpoints, confounders, model families, and adaptation rules should be prespecified to limit researcher degrees of freedom."
]:
    doc.add_paragraph(text, style="List Bullet")

add_callout(doc, "Final conclusion", "The strongest proposal is a national-to-local validation study: learn PFAS–liver and PFAS–lipid relationships in harmonized NHANES, test them without modification in an exposed Florida cohort, measure where and why transport fails, and apply transparent recalibration only after reporting the untouched external performance.", PALE_BLUE)

# References
doc.add_page_break()
doc.add_heading("References and data sources", level=1)
references = [
    ("Borghese MM et al. (2022). Individual and mixture associations of PFAS on liver function biomarkers in the Canadian Health Measures Survey.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC9472375/"),
    ("Yan Y et al. (2024). Association of exposure to PFAS with liver injury in American adults.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC11629162/"),
    ("Costello E et al. (2022). Exposure to PFAS and markers of liver injury: systematic review and meta-analysis.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC9044977/"),
    ("Haug M et al. (2023). Associations of PFAS with lipid and lipoprotein profiles.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC10541331/"),
    ("Hall SM et al. (2024). PFAS levels in paired drinking water and serum samples from Pittsboro, North Carolina.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC10529814/"),
    ("Cheng W et al. (2023). PFAS exposure and hepatic fibrosis versus steatosis: NHANES 2017–2018.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC10281433/"),
    ("Momo K et al. (2024). PFAS and NAFLD in U.S. adults, 2003–2018.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC10852365/"),
    ("Lin C-Y et al. (2010). Low-dose serum perfluorinated chemicals and liver enzymes in U.S. adults.", "https://pubmed.ncbi.nlm.nih.gov/20010922/"),
    ("Gleason JA et al. (2015). PFAS serum concentrations and biomarkers of liver function in NHANES 2007–2010.", "https://pubmed.ncbi.nlm.nih.gov/25460614/"),
    ("Jain RB, Ducatman A. (2019). PFAS and liver function biomarkers in NHANES 2011–2014.", "https://pubmed.ncbi.nlm.nih.gov/30589657/"),
    ("Jin R et al. (2020). PFAS and severity of pediatric NAFLD: an untargeted metabolomics approach.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC6944061/"),
    ("Goodrich JA et al. (2022). PFAS exposure and hepatocellular carcinoma risk in the Multiethnic Cohort.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC9468464/"),
    ("Steenland K et al. (2009). PFOA/PFOS and serum lipids among adults near a chemical plant.", "https://pubmed.ncbi.nlm.nih.gov/19846564/"),
    ("Frisbee SJ et al. (2010). PFOA/PFOS and serum lipids in children and adolescents: C8 Health Project.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC3116641/"),
    ("Nelson JW et al. (2010). Polyfluoroalkyl chemicals, cholesterol, body weight and insulin resistance in NHANES.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC2831917/"),
    ("Fitz-Simon N et al. (2013). Reductions in serum lipids with declining PFOA and PFOS.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC4724201/"),
    ("Lind PM et al. (2022). Changes in PFAS and plasma lipids over ten years in PIVUS.", "https://doi.org/10.1016/j.envres.2022.112903"),
    ("Liu B et al. (2023). PFAS exposure and blood lipid levels among adults: meta-analysis.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC10159273/"),
    ("GenX Exposure Study lipid analysis (2022). Drinking-water-associated PFAS/fluoroethers and lipid outcomes.", "https://pmc.ncbi.nlm.nih.gov/articles/PMC9450637/"),
    ("CDC/NCHS. NHANES laboratory data search and documentation.", "https://wwwn.cdc.gov/Nchs/Nhanes/Search/DataPage.aspx?Component=Laboratory"),
    ("Statistics Canada. Research Data Centre access.", "https://www.statcan.gc.ca/en/microdata/data-centres/access"),
    ("ATSDR. Findings across ten PFAS exposure-assessment sites.", "https://www.atsdr.cdc.gov/pfas/final-report/index.html"),
    ("EPA. Fifth Unregulated Contaminant Monitoring Rule Data Finder.", "https://www.epa.gov/dwucmr/fifth-unregulated-contaminant-monitoring-rule-data-finder"),
    ("USGS. PFAS concentrations in tap water throughout the United States.", "https://www.usgs.gov/data/concentrations-and-polyfluoroalkyl-substances-pfas-tapwater-collected-throughout-united-states"),
    ("Florida DEP. Drinking-water chemical data.", "https://floridadep.gov/water/source-drinking-water/content/chemical-data"),
]
for i, (citation, url) in enumerate(references, 1):
    p = doc.add_paragraph(style=None)
    p.paragraph_format.left_indent = Inches(.18)
    p.paragraph_format.first_line_indent = Inches(-.18)
    p.add_run(f"{i}. {citation} ")
    add_hyperlink(p, "Link", url)

doc.add_paragraph()
p = doc.add_paragraph()
r = p.add_run("Interpretation note: ")
r.bold = True
p.add_run("Most human studies summarized here are observational. Associations, mixture weights, and predictive performance do not by themselves demonstrate that PFAS caused the observed biomarker changes. Consult the full papers before quoting numerical estimates in a protocol or proposal.")

doc.core_properties.title = "PFAS datasets, evidence, and proposed study design"
doc.core_properties.subject = "PFAS, liver health, lipids, datasets, external validation"
doc.core_properties.author = "Prepared for PFAS proposal development"
doc.core_properties.keywords = "PFAS; NHANES; liver; lipids; datasets; external validation"
doc.save(OUT)
print(OUT)
